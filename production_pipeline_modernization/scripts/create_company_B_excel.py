import openpyxl
from openpyxl import Workbook
from datetime import date, datetime
import os

# Create output directory
output_dir = r"C:\Users\User\Desktop\Delta_lake_project\production_pipeline_modernization\test_data\company_B"
os.makedirs(output_dir, exist_ok=True)

# ===================================
# Create HR Data Excel for Company B
# ===================================
hr_wb = Workbook()

# Remove default sheet
if 'Sheet' in hr_wb.sheetnames:
    hr_wb.remove(hr_wb['Sheet'])

# Sheet 1: Employees
ws_employees = hr_wb.create_sheet("Employees", 0)
ws_employees.append(["EmployeeID", "FirstName", "LastName", "Email", "Department", "HireDate", "Salary"])
employees_data = [
    ["E011", "Thomas", "Anderson", "thomas.anderson@companyb.com", "HR", "2021-02-10", 68000],
    ["E012", "Linda", "Taylor", "linda.taylor@companyb.com", "IT", "2020-05-18", 78000],
    ["E013", "Richard", "Moore", "richard.moore@companyb.com", "Finance", "2022-01-15", 73000],
    ["E014", "Patricia", "White", "patricia.white@companyb.com", "Marketing", "2021-08-20", 65000],
    ["E015", "Charles", "Harris", "charles.harris@companyb.com", "IT", "2019-11-12", 82000],
    ["E016", "Barbara", "Martin", "barbara.martin@companyb.com", "HR", "2023-03-05", 60000],
    ["E017", "Daniel", "Thompson", "daniel.thompson@companyb.com", "Finance", "2020-09-22", 75000],
    ["E018", "Nancy", "Garcia", "nancy.garcia@companyb.com", "Marketing", "2022-06-10", 67000],
    ["E019", "Matthew", "Miller", "matthew.miller@companyb.com", "IT", "2021-04-15", 79000],
    ["E020", "Karen", "Davis", "karen.davis@companyb.com", "HR", "2020-10-08", 63000]
]
for row in employees_data:
    ws_employees.append(row)

# Sheet 2: Departments
ws_departments = hr_wb.create_sheet("Departments", 1)
ws_departments.append(["DepartmentID", "DepartmentName", "ManagerID", "Location", "Budget"])
departments_data = [
    ["D005", "HR", "E011", "Building D", 550000],
    ["D006", "IT", "E012", "Building E", 1300000],
    ["D007", "Finance", "E013", "Building D", 850000],
    ["D008", "Marketing", "E014", "Building F", 650000]
]
for row in departments_data:
    ws_departments.append(row)

# Sheet 3: Positions
ws_positions = hr_wb.create_sheet("Positions", 2)
ws_positions.append(["PositionID", "PositionTitle", "DepartmentID", "MinSalary", "MaxSalary"])
positions_data = [
    ["P009", "HR Director", "D005", 60000, 75000],
    ["P010", "Senior Developer", "D006", 75000, 95000],
    ["P011", "Finance Manager", "D007", 70000, 90000],
    ["P012", "Marketing Lead", "D008", 60000, 75000],
    ["P013", "DevOps Engineer", "D006", 70000, 90000],
    ["P014", "HR Specialist", "D005", 50000, 65000],
    ["P015", "Lead Developer", "D006", 90000, 115000],
    ["P016", "Senior Accountant", "D007", 60000, 80000]
]
for row in positions_data:
    ws_positions.append(row)

# Sheet 4: Locations
ws_locations = hr_wb.create_sheet("Locations", 3)
ws_locations.append(["LocationID", "BuildingName", "Address", "City", "State", "ZipCode", "Capacity"])
locations_data = [
    ["L004", "Building D", "321 Corporate Dr", "Austin", "TX", "78701", 250],
    ["L005", "Building E", "654 Innovation Way", "Seattle", "WA", "98101", 350],
    ["L006", "Building F", "987 Business Park", "Boston", "MA", "02101", 180]
]
for row in locations_data:
    ws_locations.append(row)

# Sheet 5: Performance
ws_performance = hr_wb.create_sheet("Performance", 4)
ws_performance.append(["ReviewID", "EmployeeID", "ReviewDate", "Rating", "Reviewer", "Comments"])
performance_data = [
    ["R009", "E011", "2025-07-10", 4.3, "E012", "Strong leadership skills"],
    ["R010", "E012", "2025-07-11", 4.9, "E011", "Exceptional technical expertise"],
    ["R011", "E013", "2025-07-12", 4.4, "E012", "Excellent financial planning"],
    ["R012", "E014", "2025-07-13", 4.5, "E011", "Innovative marketing approach"],
    ["R013", "E015", "2025-07-14", 4.8, "E012", "Outstanding development work"],
    ["R014", "E016", "2025-07-15", 4.1, "E011", "Good HR coordination"],
    ["R015", "E017", "2025-07-16", 4.4, "E013", "Reliable financial analysis"],
    ["R016", "E018", "2025-07-17", 4.6, "E014", "Creative campaign strategies"]
]
for row in performance_data:
    ws_performance.append(row)

# Save HR Data Excel
hr_file_path = os.path.join(output_dir, "hr_data.xlsx")
hr_wb.save(hr_file_path)
print(f"Created: {hr_file_path}")

# ========================================
# Create Payroll Data Excel for Company B
# ========================================
payroll_wb = Workbook()

# Remove default sheet
if 'Sheet' in payroll_wb.sheetnames:
    payroll_wb.remove(payroll_wb['Sheet'])

# Sheet 1: Salaries
ws_salaries = payroll_wb.create_sheet("Salaries", 0)
ws_salaries.append(["PayrollID", "EmployeeID", "PayPeriod", "GrossPay", "Deductions", "NetPay", "PaymentDate"])
salaries_data = [
    ["P011", "E011", "2025-10", 5666.67, 1133.33, 4533.34, "2025-10-31"],
    ["P012", "E012", "2025-10", 6500.00, 1300.00, 5200.00, "2025-10-31"],
    ["P013", "E013", "2025-10", 6083.33, 1216.67, 4866.66, "2025-10-31"],
    ["P014", "E014", "2025-10", 5416.67, 1083.33, 4333.34, "2025-10-31"],
    ["P015", "E015", "2025-10", 6833.33, 1366.67, 5466.66, "2025-10-31"],
    ["P016", "E016", "2025-10", 5000.00, 1000.00, 4000.00, "2025-10-31"],
    ["P017", "E017", "2025-10", 6250.00, 1250.00, 5000.00, "2025-10-31"],
    ["P018", "E018", "2025-10", 5583.33, 1116.67, 4466.66, "2025-10-31"],
    ["P019", "E019", "2025-10", 6583.33, 1316.67, 5266.66, "2025-10-31"],
    ["P020", "E020", "2025-10", 5250.00, 1050.00, 4200.00, "2025-10-31"]
]
for row in salaries_data:
    ws_salaries.append(row)

# Sheet 2: Bonuses
ws_bonuses = payroll_wb.create_sheet("Bonuses", 1)
ws_bonuses.append(["BonusID", "EmployeeID", "BonusType", "Amount", "BonusDate", "Reason"])
bonuses_data = [
    ["B007", "E011", "Performance", 5500, "2025-12-15", "Excellent team leadership"],
    ["B008", "E012", "Performance", 8000, "2025-12-15", "Major system upgrade"],
    ["B009", "E013", "Quarterly", 3500, "2025-09-30", "Q3 financial goals met"],
    ["B010", "E015", "Performance", 6500, "2025-12-15", "Innovation award"],
    ["B011", "E017", "Retention", 4500, "2025-11-01", "Long service bonus"],
    ["B012", "E019", "Performance", 6000, "2025-12-15", "Outstanding project delivery"]
]
for row in bonuses_data:
    ws_bonuses.append(row)

# Sheet 3: Deductions
ws_deductions = payroll_wb.create_sheet("Deductions", 2)
ws_deductions.append(["DeductionID", "EmployeeID", "DeductionType", "Amount", "DeductionDate"])
deductions_data = [
    ["D011", "E011", "Health Insurance", 275.00, "2025-10-31"],
    ["D012", "E011", "401k", 858.33, "2025-10-31"],
    ["D013", "E012", "Health Insurance", 325.00, "2025-10-31"],
    ["D014", "E012", "401k", 975.00, "2025-10-31"],
    ["D015", "E013", "Health Insurance", 275.00, "2025-10-31"],
    ["D016", "E013", "401k", 941.67, "2025-10-31"],
    ["D017", "E014", "Health Insurance", 275.00, "2025-10-31"],
    ["D018", "E014", "401k", 808.33, "2025-10-31"],
    ["D019", "E015", "Health Insurance", 375.00, "2025-10-31"],
    ["D020", "E015", "401k", 991.67, "2025-10-31"]
]
for row in deductions_data:
    ws_deductions.append(row)

# Sheet 4: Benefits
ws_benefits = payroll_wb.create_sheet("Benefits", 3)
ws_benefits.append(["BenefitID", "EmployeeID", "BenefitType", "Provider", "MonthlyPremium", "StartDate"])
benefits_data = [
    ["BN009", "E011", "Health Insurance", "Aetna", 275.00, "2021-02-10"],
    ["BN010", "E011", "Dental Insurance", "Cigna", 55.00, "2021-02-10"],
    ["BN011", "E012", "Health Insurance", "Aetna", 325.00, "2020-05-18"],
    ["BN012", "E012", "Vision Insurance", "EyeMed", 30.00, "2020-05-18"],
    ["BN013", "E013", "Health Insurance", "Aetna", 275.00, "2022-01-15"],
    ["BN014", "E014", "Health Insurance", "Aetna", 275.00, "2021-08-20"],
    ["BN015", "E015", "Health Insurance", "Aetna", 375.00, "2019-11-12"],
    ["BN016", "E015", "Life Insurance", "Prudential", 80.00, "2019-11-12"]
]
for row in benefits_data:
    ws_benefits.append(row)

# Sheet 5: Timesheets
ws_timesheets = payroll_wb.create_sheet("Timesheets", 4)
ws_timesheets.append(["TimesheetID", "EmployeeID", "WeekEnding", "RegularHours", "OvertimeHours", "TotalHours"])
timesheets_data = [
    ["T011", "E011", "2025-10-05", 40, 0, 40],
    ["T012", "E011", "2025-10-12", 40, 3, 43],
    ["T013", "E012", "2025-10-05", 40, 6, 46],
    ["T014", "E012", "2025-10-12", 40, 4, 44],
    ["T015", "E013", "2025-10-05", 40, 0, 40],
    ["T016", "E014", "2025-10-05", 40, 0, 40],
    ["T017", "E015", "2025-10-05", 40, 10, 50],
    ["T018", "E015", "2025-10-12", 40, 7, 47],
    ["T019", "E016", "2025-10-05", 40, 0, 40],
    ["T020", "E017", "2025-10-05", 40, 3, 43]
]
for row in timesheets_data:
    ws_timesheets.append(row)

# Save Payroll Data Excel
payroll_file_path = os.path.join(output_dir, "payroll_data.xlsx")
payroll_wb.save(payroll_file_path)
print(f"Created: {payroll_file_path}")

print("\nCompany B Excel files created successfully!")
print(f"HR Data: {hr_file_path}")
print(f"Payroll Data: {payroll_file_path}")
