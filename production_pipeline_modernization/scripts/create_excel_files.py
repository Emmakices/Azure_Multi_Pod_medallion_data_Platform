import openpyxl
from openpyxl import Workbook
from datetime import date, datetime
import os

# Create output directory
output_dir = r"C:\Users\User\Desktop\Delta_lake_project\production_pipeline_modernization\test_data\company_A"
os.makedirs(output_dir, exist_ok=True)

# ===================================
# Create HR Data Excel (hr_data.xlsx)
# ===================================
hr_wb = Workbook()

# Remove default sheet
if 'Sheet' in hr_wb.sheetnames:
    hr_wb.remove(hr_wb['Sheet'])

# Sheet 1: Employees
ws_employees = hr_wb.create_sheet("Employees", 0)
ws_employees.append(["EmployeeID", "FirstName", "LastName", "Email", "Department", "HireDate", "Salary"])
employees_data = [
    ["E001", "John", "Smith", "john.smith@company.com", "HR", "2020-01-15", 65000],
    ["E002", "Sarah", "Johnson", "sarah.johnson@company.com", "IT", "2019-03-22", 75000],
    ["E003", "Michael", "Williams", "michael.williams@company.com", "Finance", "2021-06-10", 70000],
    ["E004", "Emily", "Brown", "emily.brown@company.com", "Marketing", "2020-11-05", 62000],
    ["E005", "David", "Jones", "david.jones@company.com", "IT", "2018-09-18", 80000],
    ["E006", "Lisa", "Garcia", "lisa.garcia@company.com", "HR", "2022-02-28", 58000],
    ["E007", "James", "Martinez", "james.martinez@company.com", "Finance", "2019-07-14", 72000],
    ["E008", "Jennifer", "Rodriguez", "jennifer.rodriguez@company.com", "Marketing", "2021-04-20", 64000],
    ["E009", "Robert", "Lee", "robert.lee@company.com", "IT", "2020-08-11", 77000],
    ["E010", "Maria", "Gonzalez", "maria.gonzalez@company.com", "HR", "2019-12-03", 61000]
]
for row in employees_data:
    ws_employees.append(row)

# Sheet 2: Departments
ws_departments = hr_wb.create_sheet("Departments", 1)
ws_departments.append(["DepartmentID", "DepartmentName", "ManagerID", "Location", "Budget"])
departments_data = [
    ["D001", "HR", "E001", "Building A", 500000],
    ["D002", "IT", "E002", "Building B", 1200000],
    ["D003", "Finance", "E003", "Building A", 800000],
    ["D004", "Marketing", "E004", "Building C", 600000]
]
for row in departments_data:
    ws_departments.append(row)

# Sheet 3: Positions
ws_positions = hr_wb.create_sheet("Positions", 2)
ws_positions.append(["PositionID", "PositionTitle", "DepartmentID", "MinSalary", "MaxSalary"])
positions_data = [
    ["P001", "HR Manager", "D001", 55000, 70000],
    ["P002", "Software Engineer", "D002", 70000, 90000],
    ["P003", "Financial Analyst", "D003", 65000, 85000],
    ["P004", "Marketing Specialist", "D004", 55000, 70000],
    ["P005", "IT Support", "D002", 45000, 60000],
    ["P006", "HR Coordinator", "D001", 45000, 60000],
    ["P007", "Senior Developer", "D002", 85000, 110000],
    ["P008", "Accountant", "D003", 55000, 75000]
]
for row in positions_data:
    ws_positions.append(row)

# Sheet 4: Locations
ws_locations = hr_wb.create_sheet("Locations", 3)
ws_locations.append(["LocationID", "BuildingName", "Address", "City", "State", "ZipCode", "Capacity"])
locations_data = [
    ["L001", "Building A", "123 Main St", "New York", "NY", "10001", 200],
    ["L002", "Building B", "456 Tech Ave", "San Francisco", "CA", "94102", 300],
    ["L003", "Building C", "789 Market Blvd", "Chicago", "IL", "60601", 150]
]
for row in locations_data:
    ws_locations.append(row)

# Sheet 5: Performance
ws_performance = hr_wb.create_sheet("Performance", 4)
ws_performance.append(["ReviewID", "EmployeeID", "ReviewDate", "Rating", "Reviewer", "Comments"])
performance_data = [
    ["R001", "E001", "2025-06-15", 4.5, "E002", "Excellent performance"],
    ["R002", "E002", "2025-06-16", 4.8, "E001", "Outstanding technical skills"],
    ["R003", "E003", "2025-06-17", 4.2, "E002", "Good financial analysis"],
    ["R004", "E004", "2025-06-18", 4.6, "E001", "Creative marketing strategies"],
    ["R005", "E005", "2025-06-19", 4.7, "E002", "Exceptional coding abilities"],
    ["R006", "E006", "2025-06-20", 4.0, "E001", "Solid HR coordination"],
    ["R007", "E007", "2025-06-21", 4.3, "E003", "Accurate financial reporting"],
    ["R008", "E008", "2025-06-22", 4.4, "E004", "Innovative campaign ideas"]
]
for row in performance_data:
    ws_performance.append(row)

# Save HR Data Excel
hr_file_path = os.path.join(output_dir, "hr_data.xlsx")
hr_wb.save(hr_file_path)
print(f"Created: {hr_file_path}")

# ========================================
# Create Payroll Data Excel (payroll_data.xlsx)
# ========================================
payroll_wb = Workbook()

# Remove default sheet
if 'Sheet' in payroll_wb.sheetnames:
    payroll_wb.remove(payroll_wb['Sheet'])

# Sheet 1: Salaries
ws_salaries = payroll_wb.create_sheet("Salaries", 0)
ws_salaries.append(["PayrollID", "EmployeeID", "PayPeriod", "GrossPay", "Deductions", "NetPay", "PaymentDate"])
salaries_data = [
    ["P001", "E001", "2025-10", 5416.67, 1083.33, 4333.34, "2025-10-31"],
    ["P002", "E002", "2025-10", 6250.00, 1250.00, 5000.00, "2025-10-31"],
    ["P003", "E003", "2025-10", 5833.33, 1166.67, 4666.66, "2025-10-31"],
    ["P004", "E004", "2025-10", 5166.67, 1033.33, 4133.34, "2025-10-31"],
    ["P005", "E005", "2025-10", 6666.67, 1333.33, 5333.34, "2025-10-31"],
    ["P006", "E006", "2025-10", 4833.33, 966.67, 3866.66, "2025-10-31"],
    ["P007", "E007", "2025-10", 6000.00, 1200.00, 4800.00, "2025-10-31"],
    ["P008", "E008", "2025-10", 5333.33, 1066.67, 4266.66, "2025-10-31"],
    ["P009", "E009", "2025-10", 6416.67, 1283.33, 5133.34, "2025-10-31"],
    ["P010", "E010", "2025-10", 5083.33, 1016.67, 4066.66, "2025-10-31"]
]
for row in salaries_data:
    ws_salaries.append(row)

# Sheet 2: Bonuses
ws_bonuses = payroll_wb.create_sheet("Bonuses", 1)
ws_bonuses.append(["BonusID", "EmployeeID", "BonusType", "Amount", "BonusDate", "Reason"])
bonuses_data = [
    ["B001", "E001", "Performance", 5000, "2025-12-15", "Annual performance bonus"],
    ["B002", "E002", "Performance", 7500, "2025-12-15", "Outstanding project delivery"],
    ["B003", "E003", "Quarterly", 3000, "2025-09-30", "Q3 target achievement"],
    ["B004", "E005", "Performance", 6000, "2025-12-15", "Innovation award"],
    ["B005", "E007", "Retention", 4000, "2025-11-01", "5-year service bonus"],
    ["B006", "E009", "Performance", 5500, "2025-12-15", "Client satisfaction excellence"]
]
for row in bonuses_data:
    ws_bonuses.append(row)

# Sheet 3: Deductions
ws_deductions = payroll_wb.create_sheet("Deductions", 2)
ws_deductions.append(["DeductionID", "EmployeeID", "DeductionType", "Amount", "DeductionDate"])
deductions_data = [
    ["D001", "E001", "Health Insurance", 250.00, "2025-10-31"],
    ["D002", "E001", "401k", 833.33, "2025-10-31"],
    ["D003", "E002", "Health Insurance", 300.00, "2025-10-31"],
    ["D004", "E002", "401k", 950.00, "2025-10-31"],
    ["D005", "E003", "Health Insurance", 250.00, "2025-10-31"],
    ["D006", "E003", "401k", 916.67, "2025-10-31"],
    ["D007", "E004", "Health Insurance", 250.00, "2025-10-31"],
    ["D008", "E004", "401k", 783.33, "2025-10-31"],
    ["D009", "E005", "Health Insurance", 350.00, "2025-10-31"],
    ["D010", "E005", "401k", 983.33, "2025-10-31"]
]
for row in deductions_data:
    ws_deductions.append(row)

# Sheet 4: Benefits
ws_benefits = payroll_wb.create_sheet("Benefits", 3)
ws_benefits.append(["BenefitID", "EmployeeID", "BenefitType", "Provider", "MonthlyPremium", "StartDate"])
benefits_data = [
    ["BN001", "E001", "Health Insurance", "Blue Cross", 250.00, "2020-01-15"],
    ["BN002", "E001", "Dental Insurance", "Delta Dental", 50.00, "2020-01-15"],
    ["BN003", "E002", "Health Insurance", "Blue Cross", 300.00, "2019-03-22"],
    ["BN004", "E002", "Vision Insurance", "VSP", 25.00, "2019-03-22"],
    ["BN005", "E003", "Health Insurance", "Blue Cross", 250.00, "2021-06-10"],
    ["BN006", "E004", "Health Insurance", "Blue Cross", 250.00, "2020-11-05"],
    ["BN007", "E005", "Health Insurance", "Blue Cross", 350.00, "2018-09-18"],
    ["BN008", "E005", "Life Insurance", "MetLife", 75.00, "2018-09-18"]
]
for row in benefits_data:
    ws_benefits.append(row)

# Sheet 5: Timesheets
ws_timesheets = payroll_wb.create_sheet("Timesheets", 4)
ws_timesheets.append(["TimesheetID", "EmployeeID", "WeekEnding", "RegularHours", "OvertimeHours", "TotalHours"])
timesheets_data = [
    ["T001", "E001", "2025-10-05", 40, 0, 40],
    ["T002", "E001", "2025-10-12", 40, 2, 42],
    ["T003", "E002", "2025-10-05", 40, 5, 45],
    ["T004", "E002", "2025-10-12", 40, 3, 43],
    ["T005", "E003", "2025-10-05", 40, 0, 40],
    ["T006", "E004", "2025-10-05", 40, 0, 40],
    ["T007", "E005", "2025-10-05", 40, 8, 48],
    ["T008", "E005", "2025-10-12", 40, 6, 46],
    ["T009", "E006", "2025-10-05", 40, 0, 40],
    ["T010", "E007", "2025-10-05", 40, 2, 42]
]
for row in timesheets_data:
    ws_timesheets.append(row)

# Save Payroll Data Excel
payroll_file_path = os.path.join(output_dir, "payroll_data.xlsx")
payroll_wb.save(payroll_file_path)
print(f"Created: {payroll_file_path}")

print("\nExcel files created successfully!")
print(f"HR Data: {hr_file_path}")
print(f"Payroll Data: {payroll_file_path}")
